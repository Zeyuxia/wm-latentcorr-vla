from __future__ import annotations

import json
import re
import tempfile
from datetime import datetime
from pathlib import Path
from typing import Iterable

import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
from PIL import Image
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt

plt.rcParams['mathtext.fontset'] = 'stix'
plt.rcParams['font.family'] = 'STIXGeneral'

OUT_DIR = Path('/data/zhenyangfan/RoboTwin/policy/ACT_LatentCorr')
TODAY = datetime.now().strftime('%Y-%m-%d')
DOCX_PATH = OUT_DIR / f'ACT_LatentCorr_阶段性总结_当前路线版_{TODAY}.docx'

FAILURE_TABLE_PATH = OUT_DIR / 'outputs/formal_runs/stage2_multitask_failure_explore/stage2_multitask_failure_explore_new_evac_gpu0123_20260429_004924/failure_explore/failure_table.json'
TRAIN_LOG_PATH = OUT_DIR / 'outputs/logs/stage1_unified_failure_multitask_fulltable_lctok05_gpu0123_20260511_163455/train.log'
LAUNCH_INFO_PATH = OUT_DIR / 'outputs/formal_runs/stage1_unified_failure_multitask/stage1_unified_failure_multitask_fulltable_lctok05_gpu0123_20260511_163455/launch_info.txt'

OPEN_LAPTOP_EVAL_DIRS = {
    '350': OUT_DIR / 'outputs/logs/actmt350_gpu4567_pred_eval50_novideo_20260511_open_laptop',
    '400': OUT_DIR / 'outputs/logs/actmt400_gpu4567_pred_eval50_novideo_20260511_open_laptop',
    '450': OUT_DIR / 'outputs/logs/actmt450_gpu4567_pred_eval50_novideo_20260511_open_laptop',
    '600': OUT_DIR / 'outputs/logs/actmt600_gpu4567_pred_eval50_novideo_20260511_open_laptop',
}

TASK_DISPLAY = {
    'sim-open_laptop-demo_clean-50': 'open_laptop',
    'sim-pick_dual_bottles-demo_clean-50': 'pick_dual_bottles',
    'sim-put_bottles_dustbin-demo_clean-50': 'put_bottles_dustbin',
    'sim-place_burger_fries-demo_clean-50': 'place_burger_fries',
    'sim-handover_block-demo_clean-50': 'handover_block',
}

PRUNED_COUNTS = {
    'put_bottles_dustbin': 986,
    'place_burger_fries': 36,
    'handover_block': 507,
}

FEASIBILITY_METRICS = {
    'feature_dim': '18',
    'token_dim': '512',
    'linear_probe_mse': '3.391e-06',
    'mean_baseline_mse': '7.468e-06',
    'linear_probe_gain': '2.20',
    'knn_future_mse': '3.347e-06',
    'global_adjacent_pair_mse': '1.504e-05',
    'exact_duplicate_groups': '1',
    'exact_duplicate_target_mse': '4.099e-07',
}


def set_run_font(run, latin='Times New Roman', east='宋体', size=11, bold=False):
    run.font.name = latin
    run.font.size = Pt(size)
    run.bold = bold
    r = run._element
    rPr = r.get_or_add_rPr()
    rFonts = rPr.rFonts
    if rFonts is None:
        rFonts = OxmlElement('w:rFonts')
        rPr.append(rFonts)
    rFonts.set(qn('w:ascii'), latin)
    rFonts.set(qn('w:hAnsi'), latin)
    rFonts.set(qn('w:eastAsia'), east)


def set_doc_defaults(doc: Document):
    normal = doc.styles['Normal']
    normal.font.name = 'Times New Roman'
    normal.font.size = Pt(11)
    normal._element.rPr.rFonts.set(qn('w:eastAsia'), '宋体')


def add_text_paragraph(doc: Document, text: str = '', *, size=11, bold=False, align=None, latin='Times New Roman', east='宋体'):
    p = doc.add_paragraph()
    if align is not None:
        p.alignment = align
    if text:
        r = p.add_run(text)
        set_run_font(r, latin=latin, east=east, size=size, bold=bold)
    return p


def add_heading(doc: Document, text: str, level: int = 1):
    sizes = {1: 18, 2: 14, 3: 12}
    p = doc.add_paragraph()
    r = p.add_run(text)
    set_run_font(r, size=sizes.get(level, 11), bold=True)
    return p


def add_bullets(doc: Document, items: Iterable[str]):
    for item in items:
        p = doc.add_paragraph(style='List Bullet')
        r = p.add_run(item)
        set_run_font(r)


def add_numbered(doc: Document, items: Iterable[str]):
    for item in items:
        p = doc.add_paragraph(style='List Number')
        r = p.add_run(item)
        set_run_font(r)


def add_table(doc: Document, headers: list[str], rows: list[list[str]]):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = 'Table Grid'
    for i, h in enumerate(headers):
        table.rows[0].cells[i].text = h
    for cell in table.rows[0].cells:
        for p in cell.paragraphs:
            for r in p.runs:
                set_run_font(r, bold=True)
    for row in rows:
        cells = table.add_row().cells
        for i, val in enumerate(row):
            cells[i].text = val
    for row in table.rows[1:]:
        for cell in row.cells:
            for p in cell.paragraphs:
                for r in p.runs:
                    set_run_font(r)
    return table


def render_formula_png(formula_tex: str, out_path: Path, fontsize: int = 18) -> tuple[int, int]:
    fig = plt.figure(figsize=(0.01, 0.01), dpi=300)
    fig.patch.set_alpha(0.0)
    ax = fig.add_axes([0, 0, 1, 1])
    ax.axis('off')
    txt = ax.text(0.0, 0.5, f'${formula_tex}$', fontsize=fontsize, va='center', ha='left', color='black')
    fig.canvas.draw()
    bbox = txt.get_window_extent(renderer=fig.canvas.get_renderer()).expanded(1.03, 1.20)
    width_in = max(0.2, bbox.width / fig.dpi)
    height_in = max(0.2, bbox.height / fig.dpi)
    fig.set_size_inches(width_in, height_in)
    ax.set_position([0, 0, 1, 1])
    txt.set_position((0.0, 0.5))
    fig.savefig(out_path, dpi=300, transparent=True, bbox_inches='tight', pad_inches=0.02)
    plt.close(fig)
    img = Image.open(out_path)
    return img.size


def add_formula(doc: Document, formula_tex: str, tmp_dir: Path, idx: int, note: str | None = None):
    out_path = tmp_dir / f'formula_{idx:03d}.png'
    width_px, _ = render_formula_png(formula_tex, out_path)
    width_in = min(6.2, max(2.2, width_px / 220.0))
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run()
    r.add_picture(str(out_path), width=Inches(width_in))
    if note:
        pn = doc.add_paragraph()
        pn.alignment = WD_ALIGN_PARAGRAPH.CENTER
        rn = pn.add_run(note)
        set_run_font(rn, size=10)


def parse_launch_info(path: Path) -> dict[str, str]:
    info = {}
    if not path.exists():
        return info
    for line in path.read_text(errors='ignore').splitlines():
        line = line.strip()
        if not line or '=' not in line:
            continue
        k, v = line.split('=', 1)
        info[k.strip()] = v.strip()
    return info


def parse_failure_table(path: Path) -> tuple[int, dict[str, int]]:
    if not path.exists():
        return 0, {}
    data = json.loads(path.read_text())
    entries = []
    if isinstance(data, dict):
        if 'entries' in data and isinstance(data['entries'], list):
            entries = data['entries']
        elif 'units' in data and isinstance(data['units'], list):
            entries = data['units']
        elif all(isinstance(v, list) for v in data.values()):
            for task, vals in data.items():
                for e in vals:
                    if isinstance(e, dict) and 'task_name' not in e:
                        e = dict(e)
                        e['task_name'] = task
                    entries.append(e)
    elif isinstance(data, list):
        entries = data
    counts: dict[str, int] = {}
    for e in entries:
        if not isinstance(e, dict):
            continue
        task = e.get('task_name') or e.get('task') or e.get('env_name') or 'UNKNOWN'
        counts[task] = counts.get(task, 0) + 1
    return len(entries), counts


def parse_latest_epoch_summary(path: Path) -> dict[str, str]:
    if not path.exists():
        return {}
    pattern = re.compile(
        r'\[epoch\s+(\d+)\]\s+loss=([0-9.]+)\s+action=([0-9.]+)\s+cond=([0-9.]+)\s+ctoken=([0-9.]+)\s+align=([0-9.]+)\s+dyn=([0-9.]+)\s+beta=([0-9.]+)\s+lambda_cond=([0-9.]+)\s+failure_valid=([0-9.]+)\s+failure_skip=([0-9.]+)\s+time=([0-9.]+)s\s+builder=([0-9.]+)s\s+rollout=([0-9.]+)s'
    )
    last = None
    for line in path.read_text(errors='ignore').splitlines():
        m = pattern.search(line)
        if m:
            last = m.groups()
    if last is None:
        return {}
    keys = [
        'epoch', 'loss', 'action', 'cond', 'ctoken', 'align', 'dyn', 'beta',
        'lambda_cond', 'failure_valid', 'failure_skip', 'epoch_time_s',
        'builder_s', 'rollout_s'
    ]
    return {k: v for k, v in zip(keys, last)}


def parse_open_laptop_eval(dir_map: dict[str, Path]) -> list[list[str]]:
    rows = []
    for epoch, root in sorted(dir_map.items(), key=lambda x: int(x[0])):
        total_s = 0
        total_n = 0
        if not root.exists():
            rows.append([epoch, '-', '-', '-'])
            continue
        for fp in sorted(root.glob('shard_*.log')):
            s = n = None
            for line in reversed(fp.read_text(errors='ignore').splitlines()):
                if 'Success rate:' in line:
                    part = line.split('Success rate:')[1].strip().split(',')[0]
                    frac = part.split('=>')[0].strip()
                    s, n = map(int, frac.split('/'))
                    break
            if s is not None:
                total_s += s
                total_n += n
        rate = f'{100.0 * total_s / total_n:.1f}%' if total_n else '-'
        rows.append([epoch, str(total_s), str(total_n), rate])
    return rows


def main():
    launch_info = parse_launch_info(LAUNCH_INFO_PATH)
    failure_total, failure_counts = parse_failure_table(FAILURE_TABLE_PATH)
    latest_epoch = parse_latest_epoch_summary(TRAIN_LOG_PATH)
    eval_rows = parse_open_laptop_eval(OPEN_LAPTOP_EVAL_DIRS)

    config_rows = [
        ['训练主线', '基础 ACT 权重 + ACT-aligned explore + unified stage1 + pred 推理'],
        ['任务数', '5 个多任务单视角任务'],
        ['相机', 'cam_high（单视角）'],
        ['动作 horizon H', launch_info.get('act_chunk_size', '50')],
        ['condition chunk / future offset', '16'],
        ['每卡 batch', f"normal={launch_info.get('normal_batch_size', '4')}, failure={launch_info.get('failure_batch_size', '2')}"] ,
        ['全局 batch', launch_info.get('global_batch_size', '24')],
        ['基础初始化权重', Path(launch_info.get('act_init_ckpt', 'policy_epoch_2000_seed_0.ckpt')).name],
        ['EVAC checkpoint', Path(launch_info.get('evac_ckpt', 'epoch=333-step=10000.ckpt')).name],
        ['EVAC config', Path(launch_info.get('evac_config', 'train_config_robotwin_new_mixed50p12_plus_pi05_rollout.yaml')).name],
        ['failure future latent mode', launch_info.get('failure_future_latent_mode', 'rollout')],
        ['lambda_action', launch_info.get('lambda_action', '1.0')],
        ['lambda_action_conditioned', launch_info.get('lambda_action_conditioned', '0.5')],
        ['schedule_action_conditioned', launch_info.get('schedule_action_conditioned', 'true')],
        ['lambda_condition_token', launch_info.get('lambda_condition_token', '0.5')],
        ['lambda_align', launch_info.get('lambda_align', '0.0')],
        ['beta_dynamics_max', launch_info.get('beta_dynamics_max', '1.0')],
    ]

    latest_summary_lines = []
    if latest_epoch:
        latest_summary_lines = [
            f"当前 clean restart 日志中，最新已完成 epoch 为 {latest_epoch['epoch']}。",
            f"对应 loss 摘要为：total={latest_epoch['loss']}，action={latest_epoch['action']}，pred-conditioned={latest_epoch['cond']}，ctoken={latest_epoch['ctoken']}，dynamics={latest_epoch['dyn']}。",
            f"当前 warmup 已拉满，beta={latest_epoch['beta']}；conditioned loss 实际权重={latest_epoch['lambda_cond']}。",
            f"failure batch 的有效条目数约为 {latest_epoch['failure_valid']}，skip 约为 {latest_epoch['failure_skip']}。",
            f"每个 epoch 的墙钟时间约 {latest_epoch['epoch_time_s']} 秒，其中 builder 约 {latest_epoch['builder_s']} 秒，rollout 约 {latest_epoch['rollout_s']} 秒。",
        ]

    failure_rows = [[TASK_DISPLAY.get(task, task), str(cnt)] for task, cnt in sorted(failure_counts.items(), key=lambda x: x[1], reverse=True)]
    prune_rows = [[k, str(v)] for k, v in PRUNED_COUNTS.items()]
    metric_rows = [[k, v] for k, v in FEASIBILITY_METRICS.items()]

    doc = Document()
    set_doc_defaults(doc)

    add_text_paragraph(doc, 'ACT_LatentCorr 当前路线阶段总结', size=20, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
    add_text_paragraph(doc, f'生成时间：{TODAY}', size=11, align=WD_ALIGN_PARAGRAPH.CENTER)
    add_text_paragraph(doc, '本文档只描述当前 ACT_LatentCorr 的现行路线，包括 motivation、workflow、算法原理、当前实现语义、当前实验现状与当前已知问题，不讨论历史演进，不讨论备选方案，也不讨论后续实验规划。')

    add_heading(doc, '1. 当前路线的目标与动机', level=1)
    add_text_paragraph(doc, '当前路线的目标，是在保留 ACT 多任务 open-loop 基础能力的同时，把“未来状态信息”和“纠错恢复信息”真正接入训练，使策略不只会模仿成功轨迹，还能学习在偏差状态下如何回到正确轨道。')
    add_text_paragraph(doc, '这条路线的关键约束是：推理时我们只有当前观测与当前机器人状态，因此任何训练信号如果依赖未来图像、未来动作或额外的一次策略推理，都不能直接作为最终部署路径。当前路线的所有设计，都是围绕这个约束展开的。')
    add_bullets(doc, [
        '目标 1：在 normal 样本上保持多任务 ACT 的基础动作能力。',
        '目标 2：在 failure 样本上显式引入纠错未来状态监督。',
        '目标 3：让 condition 路径在训练与推理两端保持一致，不依赖训练时专属信息。',
        '目标 4：尽量少改原始 ACT 主干，把新增逻辑控制在独立目录与外接 head 中。',
    ])

    add_heading(doc, '2. 当前整体 workflow', level=1)
    add_numbered(doc, [
        '准备五任务单视角的基础 ACT 多任务权重，作为 LatentCorr 训练的统一初始化。',
        '运行 ACT-aligned explore，在关键任务相位上施加扰动，筛出模型难以自恢复的 failure unit。',
        '对 explore 结果做 task_name 回填、merge 和无效项 pruning，生成统一的 failure table。',
        '启动 unified stage1 训练。每个训练 step 混合 normal batch 与 failure batch。',
        'normal batch 使用 GT future frame 构造未来 latent；failure batch 使用错误观测与纠错动作前缀做 WM rollout 构造未来 latent。',
        '训练时通过 FutureTokenPredictor 从当前观测预测可部署的 condition token；评测和部署时默认走 pred 推理路径。',
    ])

    add_heading(doc, '3. 当前 explore 与 failure table 语义', level=1)
    add_text_paragraph(doc, '当前 explore 的作用，不是无差别制造扰动，而是筛出“基础策略自己恢复不了”的错误模式。一个 failure unit 只有在多次尝试后仍然难以恢复，才值得进入 unified stage1 的 failure batch。')
    with tempfile.TemporaryDirectory() as tmp:
        tmp_dir = Path(tmp)
        idx = 1
        add_formula(doc, r'r(u) = \frac{n_{\mathrm{recover}}(u)}{k}', tmp_dir, idx); idx += 1
        add_formula(doc, r'\mathrm{keep}(u) = \mathbf{1}[r(u) < \tau]', tmp_dir, idx); idx += 1

        add_text_paragraph(doc, '当前标准设置是 k = 4，阈值 tau = 0.5，即某个 failure unit 的恢复成功率小于 0.5 时保留。')
        add_text_paragraph(doc, f'当前新 EVAC explore 产出的 failure table 总计 {failure_total} 个条目。按任务统计如下。')
        add_table(doc, ['任务', 'failure units'], failure_rows)
        add_text_paragraph(doc, '此外，当前 clean restart 训练在加载 failure table 后，还会继续剔除无法映射到本地训练数据的无效条目。至少以下任务存在明确的 pruning。')
        add_table(doc, ['任务', '训练启动时 pruned 的无效条目数'], prune_rows)

        add_heading(doc, '4. 当前模型路径与 condition token 定义', level=1)
        add_text_paragraph(doc, '当前训练的输入是当前时刻的观测与机器人状态。为了定义当前路线，我们先给出当前 teacher latent、deployable condition token 和动作输出的数学形式。')
        add_formula(doc, r'x_t = (o_t, q_t)', tmp_dir, idx); idx += 1
        add_formula(doc, r'z_t^{\mathrm{act}} = f_{\mathrm{act}}(o_t)', tmp_dir, idx); idx += 1
        add_formula(doc, r'z_t^{\mathrm{proj}} = \Pi\!\left(z_t^{\mathrm{act}}\right)', tmp_dir, idx); idx += 1
        add_text_paragraph(doc, '这里，ACT 主干先提取当前视觉特征，再由 projector 将 ACT 视觉特征映射到和 WM latent 对齐的共享空间。')
        add_text_paragraph(doc, '训练中的 future teacher latent 分为 normal 路径与 failure 路径两种语义。')
        add_formula(doc, r'z_{t+\Delta}^{\mathrm{normal}} = E_{\mathrm{wm}}\!\left(o_{t+\Delta}^{\mathrm{gt}}\right)', tmp_dir, idx); idx += 1
        add_formula(doc, r'z_{t+\Delta}^{\mathrm{failure}} = R_{\mathrm{wm}}\!\left(o_t^{\mathrm{err}}, q_t^{\mathrm{err}}, a_{\mathrm{corr},t:t+\Delta-1}\right)', tmp_dir, idx); idx += 1
        add_text_paragraph(doc, '也就是说，normal batch 的 future latent 来自 GT future frame 的 VAE 编码；failure batch 的 future latent 来自当前错误状态加纠错动作前缀的 rollout 结果。')
        add_formula(doc, r'\tilde{z}_{t+\Delta}^{\mathrm{teacher}} = A\!\left(z_{t+\Delta}^{\mathrm{teacher}}\right)', tmp_dir, idx); idx += 1
        add_formula(doc, r'c_t^{\star} = P\!\left(\mathrm{pool}\!\left(\tilde{z}_{t+\Delta}^{\mathrm{teacher}}\right)\right)', tmp_dir, idx); idx += 1
        add_text_paragraph(doc, '当前路线中，teacher future token 仍然存在，但它只作为监督目标存在，不直接作为部署时的输入。部署时真正使用的是由当前输入直接预测出的 token。')
        add_formula(doc, r'h_t = [\mathrm{pool}(z_t^{\mathrm{proj}});\, q_t]', tmp_dir, idx); idx += 1
        add_formula(doc, r'\hat{c}_t = g_{\phi}(h_t)', tmp_dir, idx); idx += 1
        add_text_paragraph(doc, '其中 g_φ 就是当前代码中的 FutureTokenPredictor。它读取当前 projector latent 的全局池化结果和当前 qpos，输出一个 ACT hidden space 中的 condition token。')

        add_heading(doc, '5. 当前训练目标', level=1)
        add_text_paragraph(doc, '当前训练里真正参与优化的核心量有四个：base 动作损失、pred-conditioned 动作损失、condition token 对齐损失以及 dynamics 损失。')
        add_formula(doc, r'\hat{a}_{t:t+H-1}^{\mathrm{base}} = \pi_{\mathrm{ACT}}(x_t)', tmp_dir, idx); idx += 1
        add_formula(doc, r'\hat{a}_{t:t+H-1}^{\mathrm{pred}} = \pi_{\mathrm{ACT}}(x_t;\, \hat{c}_t)', tmp_dir, idx); idx += 1
        add_formula(doc, r'\mathcal{L}_{\mathrm{base}} = \ell\!\left(\hat{a}_{t:t+H-1}^{\mathrm{base}},\, a_{t:t+H-1}\right)', tmp_dir, idx); idx += 1
        add_formula(doc, r'\mathcal{L}_{\mathrm{pred}} = \ell\!\left(\hat{a}_{t:t+H-1}^{\mathrm{pred}},\, a_{t:t+H-1}\right)', tmp_dir, idx); idx += 1
        add_formula(doc, r'\mathcal{L}_{\mathrm{tok}} = \left\Vert \hat{c}_t - c_t^{\star} \right\Vert_2^2', tmp_dir, idx); idx += 1
        add_formula(doc, r'\hat{z}_{t+\Delta} = F\!\left(z_t^{\mathrm{proj}}, a_{t:t+\Delta-1}\right)', tmp_dir, idx); idx += 1
        add_formula(doc, r'\mathcal{L}_{\mathrm{dyn}} = \left\Vert \hat{z}_{t+\Delta} - \tilde{z}_{t+\Delta}^{\mathrm{teacher}} \right\Vert_2^2', tmp_dir, idx); idx += 1
        add_formula(doc, r'\mathcal{L} = \lambda_a \mathcal{L}_{\mathrm{base}} + \lambda_c(t) \mathcal{L}_{\mathrm{pred}} + \beta(t)\lambda_{\mathrm{tok}}\mathcal{L}_{\mathrm{tok}} + \beta(t)\mathcal{L}_{\mathrm{dyn}}', tmp_dir, idx); idx += 1
        add_text_paragraph(doc, '当前实现里，conditioned action loss 会随 warmup 渐进增加；token loss 与 dynamics loss 也都共享同一个 warmup 系数。这样做的原因是训练初期 predictor 还不稳定，不能过早让 condition 分支主导动作学习。')
        add_text_paragraph(doc, '另外，虽然日志中仍然会打印 align loss，但当前 clean restart 训练里 lambda_align=0.0，因此 align 项不参与优化；wm_action_current、wm_action_future、bridge_future 这些历史项也都设为 0。')

        add_heading(doc, '6. 当前 normal batch 与 failure batch 的具体语义', level=1)
        add_bullets(doc, [
            'normal batch：保持基础多任务行为克隆能力，future latent 来自 GT future frame 的 VAE encode。',
            'failure batch：在线读取 failure table 条目，经 ACT-aligned correction builder 构造当前错误图像、当前错误状态和纠错动作 chunk。',
            'failure batch 的 future latent 不是“错误图像自己的 latent”，而是“错误图像 + 纠错动作前缀”经过 WM rollout 后的 future latent。',
            '因此 failure 路径学习到的不是“识别错误长什么样”，而是“从当前错误状态执行纠错后应该到达哪个未来状态”。',
        ])

        add_heading(doc, '7. 当前推理路径', level=1)
        add_text_paragraph(doc, '当前评测和部署默认使用 pred 路径。它是一张图中的单次 forward，不需要先做一次未来动作推理、再做一次当前动作推理。')
        add_formula(doc, r'x_t \rightarrow z_t^{\mathrm{act}} \rightarrow z_t^{\mathrm{proj}} \rightarrow \hat{c}_t', tmp_dir, idx); idx += 1
        add_formula(doc, r'(x_t,\, \hat{c}_t) \rightarrow \hat{a}_{t:t+H-1}^{\mathrm{pred}}', tmp_dir, idx); idx += 1
        add_text_paragraph(doc, '也就是说，新的 condition token 是当前观测分支上的一个 predictor head 直接给出的，而不是通过第二次 ACT 解码额外推出来的。')

        add_heading(doc, '8. 当前 condition token 路线的已完成可行性验证', level=1)
        add_text_paragraph(doc, '当前路线的一个核心问题，是当前输入到 future token 的映射是否严重 one-to-many。现有验证并不是在 raw pixel 上做，而是用与真实训练路径一致的特征：pool(projector latent) 与 qpos 的拼接。')
        add_formula(doc, r'\phi(x_t) = [\mathrm{pool}(z_t^{\mathrm{proj}});\, q_t]', tmp_dir, idx); idx += 1
        add_formula(doc, r'\hat{c}_t = W\phi(x_t) + b', tmp_dir, idx); idx += 1
        add_formula(doc, r'\hat{c}_t = \mathrm{kNN}\!\left(\phi(x_t)\right)', tmp_dir, idx); idx += 1
        add_formula(doc, r'\hat{c}_t = \bar{c}', tmp_dir, idx); idx += 1
        add_text_paragraph(doc, '线性 probe、kNN 回归与均值基线的对比结果如下。')
        add_table(doc, ['指标', '数值'], metric_rows)
        add_text_paragraph(doc, '这些结果说明：当前输入里确实包含可恢复的 future token 结构；至少在当前数据定义下，并没有看到严重到足以直接否掉 predictor 路线的 one-to-many 崩坏。')

    add_heading(doc, '9. 当前训练配置与现状', level=1)
    add_table(doc, ['项目', '当前设置'], config_rows)
    for line in latest_summary_lines:
        add_text_paragraph(doc, line)
    add_text_paragraph(doc, '当前已有一轮较干净的 open_laptop 单任务评测，使用 pred 推理、50 seeds、no video。结果如下。')
    add_table(doc, ['checkpoint epoch', '成功数', '总 trial', '成功率'], eval_rows)

    add_heading(doc, '10. 当前已知问题', level=1)
    add_numbered(doc, [
        '多任务收益不稳定。当前路线在部分任务上能给出正向信号，但在其他任务上提升不明显，甚至存在副作用，说明多任务层面的收益释放还不稳定。',
        'checkpoint 表现明显非单调。某些较早 epoch 的结果可能优于更晚 epoch，这意味着继续训练并不自动等于效果更好。',
        'failure table 分布不均衡。不同任务的 failure unit 数量差异很大，而且训练启动时还会进一步 pruning，导致各任务实际看到的 failure 覆盖并不一致。',
        'failure batch 的在线构造开销较大。当前 correction builder 和 rollout 都在训练时动态执行，因此 failure 路径天然比纯 normal batch 重。',
        '当前 FutureTokenPredictor 使用 global average pooling，空间细节在进入 token 预测头之前已经被压缩，细粒度偏差是否会因此丢失，目前仍然是一个客观风险。',
        'ctoken 在后期日志里常显示为 0.0000，但这大概率只是打印精度不够，因此目前难以仅靠日志判断 token 分支是否仍在持续学习。',
        'pred 路径虽然在理论上已经闭合，但它的工程链路仍比 base 路径复杂，历史上也出现过评测卡死、依赖初始化慢等现象。',
        'EVAC 权重的变化会改变 explore 的质量和 failure table 的定义，因此 failure 数据并不是一次生成后永远通用。',
    ])

    add_heading(doc, '11. 关键文件与路径', level=1)
    path_rows = [
        ['统一训练入口', str(OUT_DIR / 'train_stage1_unified_failure_multitask_latent.py')],
        ['模型主体', str(OUT_DIR / 'latent_policy.py')],
        ['模块定义', str(OUT_DIR / 'latent_modules.py')],
        ['部署入口', str(OUT_DIR / 'deploy_policy.py')],
        ['当前 launcher', str(OUT_DIR / 'launch_stage1_unified_failure_multitask_fulltable_lctok05_gpu0123.sh')],
        ['当前训练日志', str(TRAIN_LOG_PATH)],
        ['当前 failure table', str(FAILURE_TABLE_PATH)],
    ]
    add_table(doc, ['名称', '路径'], path_rows)

    add_heading(doc, '12. 总结', level=1)
    add_text_paragraph(doc, '当前 ACT_LatentCorr 的路线可以概括为：在五任务单视角基础 ACT 权重之上，先通过 ACT-aligned explore 筛出真正难以自恢复的 failure 模式，再通过 normal batch 与 failure batch 的 unified stage1 训练，把基础动作学习、future latent 监督、deployable condition token 预测以及 failure 恢复监督统一到同一个训练框架中。')
    add_text_paragraph(doc, '当前路线最关键的定义有两点。第一，推理时走 pred 路径，condition token 由当前观测直接预测；第二，failure batch 的 future latent 使用 rollout 语义，而不是简单的错误图像 encode。当前已知问题主要集中在 failure 数据分布、训练动态稳定性、failure 路径开销和 token 头的信息瓶颈上。')

    doc.save(DOCX_PATH)
    print(f'DOCX_PATH={DOCX_PATH}')


if __name__ == '__main__':
    main()
