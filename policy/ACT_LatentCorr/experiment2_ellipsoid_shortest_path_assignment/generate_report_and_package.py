from __future__ import annotations

import csv
import re
import shutil
import textwrap
import zipfile
from pathlib import Path

import matplotlib.pyplot as plt
from docx import Document
from docx.oxml.ns import qn
from docx.shared import Inches, Pt
from openpyxl import Workbook


ASSIGNMENT_DIR = Path("/data/zhenyangfan/RoboTwin/policy/ACT_LatentCorr/experiment2_ellipsoid_shortest_path_assignment")
OUTPUTS_DIR = ASSIGNMENT_DIR / "outputs"
TABLES_DIR = OUTPUTS_DIR / "tables"
FIGURES_DIR = OUTPUTS_DIR / "figures"
LOGS_DIR = OUTPUTS_DIR / "logs"
REPORT_ASSETS_DIR = ASSIGNMENT_DIR / "report_assets"
REPORT_ASSETS_DIR.mkdir(parents=True, exist_ok=True)

SUMMARY_TXT = LOGS_DIR / "results_summary.txt"
RUN_LOG = LOGS_DIR / "octave_run.log"
REPORT_PATH = ASSIGNMENT_DIR / "实验项目2_椭球面上两点之间的最短距离_实验报告_MATLAB版.docx"
ZIP_PATH = ASSIGNMENT_DIR.with_suffix(".zip")


def read_csv_rows(path: Path) -> list[list[str]]:
    with path.open("r", encoding="utf-8") as f:
        return list(csv.reader(f))


def convert_csv_to_xlsx(csv_path: Path, xlsx_path: Path) -> None:
    rows = read_csv_rows(csv_path)
    wb = Workbook()
    ws = wb.active
    ws.title = xlsx_path.stem

    ws.append(["节点（序号）", "X", "Y", "Z", "与上一个节点的直线距离"])
    for row in rows[1:]:
        out_row = [row[0]]
        for value in row[1:4]:
            out_row.append(float(value))
        out_row.append(row[4] if row[4] == "\\" else float(row[4]))
        ws.append(out_row)

    widths = {"A": 16, "B": 16, "C": 16, "D": 16, "E": 20}
    for col, width in widths.items():
        ws.column_dimensions[col].width = width

    wb.save(xlsx_path)


def make_log_screenshot(log_path: Path, out_path: Path) -> None:
    lines = log_path.read_text(encoding="utf-8").splitlines()
    focus_lines = lines[-18:] if len(lines) > 18 else lines
    text = "\n".join(focus_lines)

    fig = plt.figure(figsize=(13, 6), dpi=160)
    ax = fig.add_subplot(111)
    ax.axis("off")
    ax.text(
        0.01,
        0.99,
        text,
        va="top",
        ha="left",
        family="monospace",
        fontsize=10,
    )
    fig.tight_layout()
    fig.savefig(out_path, bbox_inches="tight")
    plt.close(fig)


def parse_metrics(summary_text: str) -> dict[str, dict[str, str]]:
    pattern = re.compile(
        r"(P\d -> P\d) 路径结果\s+"
        r"测地线最短距离：([0-9.]+)\s+"
        r"分段直线长度和：([0-9.]+)\s+"
        r"节点总数：([0-9]+)\s+"
        r"逆解迭代次数：([0-9]+)",
        re.MULTILINE,
    )
    metrics: dict[str, dict[str, str]] = {}
    for match in pattern.finditer(summary_text):
        metrics[match.group(1)] = {
            "distance": match.group(2),
            "chord_sum": match.group(3),
            "node_count": match.group(4),
            "iterations": match.group(5),
        }
    return metrics


def set_doc_font(document: Document) -> None:
    normal = document.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(10.5)
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")


def add_code_block(document: Document, title: str, code_text: str) -> None:
    document.add_heading(title, level=2)
    p = document.add_paragraph()
    run = p.add_run(code_text)
    run.font.name = "Consolas"
    run.font.size = Pt(8)
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Consolas")


def add_table_preview(document: Document, rows: list[list[str]], title: str, max_rows: int = 8) -> None:
    document.add_paragraph(title)
    preview = rows[: max_rows + 1]
    if preview:
        preview[0] = ["节点（序号）", "X", "Y", "Z", "与上一个节点的直线距离"]
    table = document.add_table(rows=len(preview), cols=len(preview[0]))
    table.style = "Table Grid"
    for i, row in enumerate(preview):
        for j, value in enumerate(row):
            table.cell(i, j).text = str(value)


def generate_report() -> None:
    summary_text = SUMMARY_TXT.read_text(encoding="utf-8")
    metrics = parse_metrics(summary_text)

    p1p2_rows = read_csv_rows(TABLES_DIR / "P1_P2route.csv")
    p2p3_rows = read_csv_rows(TABLES_DIR / "P2_P3route.csv")

    document = Document()
    set_doc_font(document)

    section = document.sections[0]
    section.top_margin = Inches(0.8)
    section.bottom_margin = Inches(0.8)
    section.left_margin = Inches(1.0)
    section.right_margin = Inches(1.0)

    document.add_heading("实验项目2：椭球面上两点之间的最短距离实验报告", level=0)
    document.add_paragraph("实现环境：MATLAB 脚本实现，当前版本已按 MATLAB 软件运行方式完成兼容性调整。")
    document.add_paragraph("验证说明：本机未安装商业版 MATLAB，因此本文使用 GNU Octave 对同一套 .m 程序做回归验证；代码主体采用 MATLAB 写法，并尽量避免 Octave 专属语法。")

    document.add_heading("一、问题分析与数学建模", level=1)
    document.add_paragraph(
        "题目给定的曲面是旋转椭球面，其方程可写为 x^2/a^2 + y^2/a^2 + z^2/b^2 = 1，其中 a = 6000，b = 5000。"
        "由于目标是在曲面上寻找两点之间的最短路线，本质上属于椭球面测地线问题。"
        "如果直接在三维空间中连线，只能得到弦长，无法保证路径始终位于椭球面上，因此必须采用曲面上的测地线模型。"
    )
    document.add_paragraph(
        "本实验将题目中的三维坐标点先转换为椭球上的大地纬度和经度，再利用 Vincenty 椭球测地线公式完成两类计算："
        "一是逆解，即已知起终点求测地线长度和起始方位角；二是正解，即已知起点、初始方位和走过的弧长求新的路径点。"
        "这样既能得到较稳定的最短距离估计，也能方便地按等弧长采样生成不少于 50 个节点的路径序列。"
    )
    document.add_paragraph(
        "建模过程中用到的关键转换关系为：先根据题目中的 x、y 计算 z，之后由 tan(phi) = z*a^2 / (p*b^2)、p = sqrt(x^2+y^2) 求大地纬度，"
        "经度由 atan2(y, x) 得到。测地线长度由 Vincenty 逆解公式迭代得到，路径节点则通过 Vincenty 正解公式沿最短测地线均匀采样。"
    )

    document.add_heading("二、算法设计", level=1)
    document.add_paragraph(
        "算法步骤如下：\n"
        "1. 根据题目给定的 x、y 和椭球参数计算三个点的 z 坐标，得到 P1、P2、P3 的完整三维坐标。\n"
        "2. 将三维直角坐标转换为大地纬度、经度。\n"
        "3. 对 P1-P2、P2-P3 两组点分别调用 Vincenty 逆解，求得最短测地线长度和起始方位角。\n"
        "4. 沿测地线总长度做等弧长离散，本实验每条路径取 100 个内部节点，共 102 个节点。\n"
        "5. 对每个采样位置调用 Vincenty 正解恢复三维坐标，并计算与上一个节点之间的直线距离。\n"
        "6. 将结果保存为 CSV/XLSX 表格，同时在椭球面上绘制路径图并输出汇总说明。"
    )
    document.add_paragraph(
        "该算法的优点是：模型明确、精度较高、计算稳定，且节点数量可自由调节。与单纯的随机投点或粗网格搜索相比，"
        "它直接利用了旋转椭球面的几何结构，因此更适合本题。"
    )

    document.add_heading("三、运行结果", level=1)
    document.add_paragraph(
        f"P1 到 P2 的测地线最短距离为 {metrics['P1 -> P2']['distance']}，"
        f"离散折线长度和为 {metrics['P1 -> P2']['chord_sum']}，"
        f"共输出 {metrics['P1 -> P2']['node_count']} 个节点。"
    )
    document.add_paragraph(
        f"P2 到 P3 的测地线最短距离为 {metrics['P2 -> P3']['distance']}，"
        f"离散折线长度和为 {metrics['P2 -> P3']['chord_sum']}，"
        f"共输出 {metrics['P2 -> P3']['node_count']} 个节点。"
    )
    document.add_paragraph("图 1、图 2 分别给出两条最短路径的单独示意图，图 3 为两条路径在同一椭球面上的整体效果图。")
    document.add_picture(str(FIGURES_DIR / "route_P1_P2.png"), width=Inches(5.8))
    document.add_paragraph("图1 P1 到 P2 的椭球面最短路径图")
    document.add_picture(str(FIGURES_DIR / "route_P2_P3.png"), width=Inches(5.8))
    document.add_paragraph("图2 P2 到 P3 的椭球面最短路径图")
    document.add_picture(str(FIGURES_DIR / "routes_overview.png"), width=Inches(5.8))
    document.add_paragraph("图3 两条最短路径的整体效果图")

    document.add_heading("四、运行截图与结果表预览", level=1)
    document.add_paragraph("下面给出 Octave 执行脚本时的日志截图，以及两张结果表的前若干行预览。")
    document.add_picture(str(REPORT_ASSETS_DIR / "octave_log_screenshot.png"), width=Inches(6.0))
    document.add_paragraph("图4 程序运行日志截图")
    add_table_preview(document, p1p2_rows, "表1 P1_P2route.xlsx 前 8 行预览")
    add_table_preview(document, p2p3_rows, "表2 P2_P3route.xlsx 前 8 行预览")

    document.add_heading("五、程序源代码", level=1)
    document.add_paragraph("以下附上本实验主要 MATLAB 源代码文本。完整源码已随作业文件夹一并打包。")
    code_files = [ASSIGNMENT_DIR / "run_all.m"] + sorted((ASSIGNMENT_DIR / "src").glob("*.m"))
    for code_file in code_files:
        add_code_block(document, f"源代码：{code_file.name}", code_file.read_text(encoding="utf-8"))

    document.add_page_break()
    document.add_heading("六、其他说明", level=1)
    document.add_paragraph(
        "1. 本实验已按题目要求生成两个结果表文件：P1_P2route.xlsx 与 P2_P3route.xlsx。\n"
        "2. 目前提交的主程序已经按 MATLAB 方式整理：在 MATLAB 中可直接运行 run_all.m，若本机支持 xlswrite，则会直接输出 xlsx；"
        "本次为了完成本机验证，额外保留了 CSV 转 XLSX 的辅助脚本。\n"
        "3. 与上一个节点的直线距离采用相邻离散节点的三维欧氏距离计算，因此其总和会与连续测地线长度存在极小差异，这是离散采样造成的正常现象。\n"
        "4. 每条路径共给出 102 个节点，已经满足题目“不少于 50 个节点”的要求。"
    )

    document.save(REPORT_PATH)


def package_submission() -> None:
    if ZIP_PATH.exists():
        ZIP_PATH.unlink()

    with zipfile.ZipFile(ZIP_PATH, "w", zipfile.ZIP_DEFLATED) as zf:
        for path in sorted(ASSIGNMENT_DIR.rglob("*")):
            if path.is_file():
                zf.write(path, path.relative_to(ASSIGNMENT_DIR.parent))


def main() -> None:
    convert_csv_to_xlsx(TABLES_DIR / "P1_P2route.csv", ASSIGNMENT_DIR / "P1_P2route.xlsx")
    convert_csv_to_xlsx(TABLES_DIR / "P2_P3route.csv", ASSIGNMENT_DIR / "P2_P3route.xlsx")

    shutil.copy2(SUMMARY_TXT, ASSIGNMENT_DIR / "results_summary.txt")
    shutil.copy2(
        Path("/data/zhenyangfan/RoboTwin/policy/ACT_LatentCorr/实验项目2--椭球面上两点之间的最短距离--2026.03.01更新.docx"),
        ASSIGNMENT_DIR / "原始题目.docx",
    )

    make_log_screenshot(RUN_LOG, REPORT_ASSETS_DIR / "octave_log_screenshot.png")
    generate_report()
    package_submission()


if __name__ == "__main__":
    main()
